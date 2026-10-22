import io
import os
import re
import json
import base64
import urllib.parse
from datetime import datetime, timedelta
import feedparser
import pandas as pd
import requests
from bs4 import BeautifulSoup
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from PIL import Image, ImageDraw
import streamlit as st
from apscheduler.schedulers.background import BackgroundScheduler
from apscheduler.triggers.cron import CronTrigger

# ==========================================
# 1. KONFIGURASI HALAMAN & DIREKTORI
# ==========================================
st.set_page_config(
    page_title="Media & Social Intelligence Monitor",
    page_icon="📡",
    layout="wide"
)

EXPORT_DIR = "exports"
CONFIG_FILE = "scheduler_config.json"
os.makedirs(EXPORT_DIR, exist_ok=True)

BULAN_ID = {
    1: "Januari", 2: "Februari", 3: "Maret", 4: "April", 5: "Mei", 6: "Juni",
    7: "Juli", 8: "Agustus", 9: "September", 10: "Oktober", 11: "November", 12: "Desember"
}

# ==========================================
# 2. HELPER RESOLVER URL ASLI DARI GOOGLE
# ==========================================
def resolve_clean_url(google_url):
    """
    Mengurai URL redirect Google News RSS agar menjadi URL asli penerbit berita.
    """
    if "news.google.com" not in google_url:
        return google_url
    
    try:
        # Coba ekstrak jika URL terenkripsi base64 di dalam path artikel
        match = re.search(r'articles/([^?]+)', google_url)
        if match:
            encoded_str = match.group(1)
            # Menambahkan padding jika perlu
            padded = encoded_str + '=' * (-len(encoded_str) % 4)
            decoded = base64.urlsafe_b64decode(padded.encode('utf-8'))
            # Cari format URL HTTP di dalam byte decoded
            urls = re.findall(r'https?://[^\s<>"\'\x00-\x1f]+', decoded.decode('latin1', errors='ignore'))
            if urls:
                return urls[0]
    except Exception:
        pass

    # Fallback cepat via HTTP HEAD (timeout 1.5 detik agar tidak membebani memori)
    try:
        headers = {"User-Agent": "Mozilla/5.0"}
        resp = requests.head(google_url, headers=headers, allow_redirects=True, timeout=1.5)
        if resp.url and "news.google.com" not in resp.url:
            return resp.url
    except Exception:
        pass

    return google_url

# ==========================================
# 3. ADVANCED QUERY BUILDER & POST-FILTERING
# ==========================================
def calculate_date_range_label(time_val):
    today = datetime.now()
    if time_val.endswith("h"):
        hours = int(time_val.replace("h", ""))
        start_time = today - timedelta(hours=hours)
        return f"{start_time.strftime('%H:%M')} – {today.strftime('%H:%M WIB')} ({today.day} {BULAN_ID[today.month]} {today.year})"
    
    days_map = {"1d": 1, "2d": 2, "3d": 3, "7d": 7, "30d": 30}
    if time_val in days_map:
        start_date = today - timedelta(days=days_map[time_val])
        return f"{start_date.day} {BULAN_ID[start_date.month]} {start_date.year} – {today.day} {BULAN_ID[today.month]} {today.year}"
    return f"Hingga {today.day} {BULAN_ID[today.month]} {today.year} (Semua Periode)"

def parse_keywords_list(raw_keywords):
    tokens = re.split(r'[,;\n]+', raw_keywords)
    cleaned = []
    for t in tokens:
        item = t.strip()
        if item:
            cleaned.append(item)
    return cleaned

def build_advanced_query(entity_name, raw_keywords, match_operator="AND", match_in_title=False, exclude_gov=True, custom_exclude_domains=""):
    """
    Menyusun query pencarian ketat untuk Google News.
    """
    entity_clean = entity_name.strip()
    entity_token = f'"{entity_clean}"' if " " in entity_clean else entity_clean
    keyword_list = parse_keywords_list(raw_keywords)

    formatted_keywords = []
    for kw in keyword_list:
        if " " in kw and not (kw.startswith('"') and kw.endswith('"')):
            formatted_keywords.append(f'"{kw}"')
        else:
            formatted_keywords.append(kw)

    if formatted_keywords:
        if match_operator == "AND":
            # Semua kata kunci wajib ada
            kw_part = " AND ".join(formatted_keywords)
            base_query = f"{entity_token} AND {kw_part}"
        else:
            # Salah satu kata kunci cukup
            kw_part = f"({' OR '.join(formatted_keywords)})"
            base_query = f"{entity_token} AND {kw_part}"
    else:
        base_query = entity_token

    if match_in_title:
        base_query = f"allintitle: {base_query}"

    exclude_clauses = []
    if exclude_gov:
        exclude_clauses.append("-site:go.id")

    if custom_exclude_domains.strip():
        for d in re.split(r'[,;\s]+', custom_exclude_domains):
            d_clean = d.strip().replace("https://", "").replace("http://", "").replace("/", "")
            if d_clean:
                exclude_clauses.append(f"-site:{d_clean}")

    if exclude_clauses:
        base_query = f"{base_query} {' '.join(exclude_clauses)}"

    return base_query, keyword_list

# ==========================================
# 4. GENERATOR KARTU VISUAL & SUMMARY
# ==========================================
def generate_press_capture_card(channel_type, source, title, published, snippet, sentiment):
    width, height = 800, 520
    bg_color = (255, 255, 255)
    
    if sentiment == "Negatif":
        accent_color, badge_bg = (220, 38, 38), (254, 226, 226)
    elif sentiment == "Positif":
        accent_color, badge_bg = (22, 163, 74), (220, 252, 231)
    else:
        accent_color, badge_bg = (202, 138, 4), (254, 243, 199)

    img = Image.new("RGB", (width, height), bg_color)
    draw = ImageDraw.Draw(img)

    draw.rectangle([(0, 0), (width - 1, height - 1)], outline=(203, 213, 225), width=3)
    draw.rectangle([(0, 0), (width, 55)], fill=(241, 245, 249))
    draw.rectangle([(0, 0), (10, height)], fill=accent_color)

    tag = "[MEDIA SOSIAL]" if channel_type == "medsos" else "[PORTAL BERITA]"
    draw.text((25, 18), f"{tag}  •  {source.upper()}", fill=(30, 41, 59))

    draw.rectangle([(width - 150, 12), (width - 25, 42)], fill=badge_bg)
    draw.text((width - 135, 18), sentiment.upper(), fill=accent_color)
    draw.text((25, 70), f"Waktu Rilis / Temuan: {published}", fill=(100, 116, 139))

    def wrap_text(text, limit=46):
        lines, cur = [], []
        for w in text.split():
            if sum(len(x) + 1 for x in cur) + len(w) <= limit:
                cur.append(w)
            else:
                lines.append(" ".join(cur))
                cur = [w]
        if cur:
            lines.append(" ".join(cur))
        return lines

    y = 105
    draw.text((25, y), "HEADLINE / JUDUL:", fill=(71, 85, 105))
    y += 24
    for line in wrap_text(title, limit=42)[:3]:
        draw.text((25, y), line, fill=(15, 23, 42))
        y += 26

    y += 12
    draw.line([(25, y), (width - 25, y)], fill=(226, 232, 240), width=1)
    y += 18

    draw.text((25, y), "DOKUMENTASI KUTIPAN KONTEN:", fill=(71, 85, 105))
    y += 24
    for line in wrap_text(snippet, limit=48)[:5]:
        draw.text((25, y), line, fill=(51, 65, 85))
        y += 24

    draw.rectangle([(0, height - 38), (width, height)], fill=(248, 250, 252))
    draw.text((25, height - 26), "VERIFIED BY STRATEGIC MEDIA AUDIT ENGINE", fill=(148, 163, 184))

    buf = io.BytesIO()
    img.save(buf, format="PNG")
    buf.seek(0)
    return buf.getvalue()

def extract_bullet_summary(title, text, min_bullets=2, max_bullets=3):
    cleaned = re.sub(r'<[^<]+?>', '', text).strip()
    sentences = re.split(r'(?<=[.!?])\s+', cleaned)
    bullets = [s.strip() for s in sentences if len(s.strip()) > 25]

    if len(bullets) < 2:
        bullets = [f"Fokus Utama: {title.strip()}"]
        if len(cleaned) > 20:
            bullets.append(f"Substansi Pemberitaan: {cleaned[:140]}...")
        else:
            bullets.append("Publikasi memuat perkembangan operasional dan respons kebijakan institusi.")
            
    return bullets[:max_bullets]

# ==========================================
# 5. ENGINE PEMINDAIAN DENGAN POST-FILTERING
# ==========================================
def fetch_media_and_social(entity_name, raw_keywords, time_window, max_results, include_news, include_social, include_small_media, match_operator, match_in_title, exclude_gov, custom_exclude):
    advanced_q, keyword_list = build_advanced_query(
        entity_name, raw_keywords, match_operator=match_operator, 
        match_in_title=match_in_title, exclude_gov=exclude_gov, custom_exclude_domains=custom_exclude
    )

    search_targets = []
    if include_news:
        search_targets.append(("news", advanced_q, "online_news"))
    if include_small_media:
        search_targets.append(("web", advanced_q, "local_web"))
    if include_social:
        social_dork = "(site:instagram.com OR site:facebook.com OR site:threads.net OR site:twitter.com OR site:x.com OR site:tiktok.com)"
        search_targets.append(("social", f"{advanced_q} {social_dork}", "medsos"))

    negative_lexicon = [
        "korupsi", "rugi", "bobol", "fraud", "kasus", "gagal", "demo", "kecewa",
        "macet", "krisis", "tuntutan", "phk", "ancam", "turun", "pidana", "gugat", 
        "disorot", "hukum", "sita", "penipuan", "keluhan", "masalah", "denda", "skandal", "viral", "hujat", "terbakar"
    ]
    positive_lexicon = [
        "laba", "untung", "tumbuh", "prestasi", "penghargaan", "inovasi", "sukses",
        "dividen", "ekspansi", "naik", "resmikan", "unggul", "mitra", "positif", 
        "rekor", "terbaik", "kinerja positif", "apresiasi", "capai", "solutif", "bangga", "lestari", "sosialisasi", "csr"
    ]

    items = []
    seen_links = set()

    for target_type, query_str, channel_type in search_targets:
        final_q = f"{query_str} when:{time_window}" if time_window != "all" else query_str
        encoded_q = urllib.parse.quote(final_q)
        rss_url = f"https://news.google.com/rss/search?q={encoded_q}&hl=id&gl=ID&ceid=ID:id"
        feed = feedparser.parse(rss_url)

        for entry in feed.entries:
            raw_link = entry.link
            title = entry.title
            published = entry.get("published", "Terkini")
            source = "Media Nasional"
            clean_title = title
            if " - " in title:
                parts = title.rsplit(" - ", 1)
                clean_title, source = parts[0], parts[1]

            raw_summary = entry.get("summary", "")
            soup = BeautifulSoup(raw_summary, "html.parser")
            snippet = soup.get_text()

            # ------------------------------------------
            # TAHAP POST-FILTERING VERIFIKASI SUBSTANSI
            # Mencegah berita dividen/kredit macet lolos jika tidak ada kata kunci CSR/Taspen
            # ------------------------------------------
            content_combined = f"{clean_title} {snippet}".lower()
            entity_lower = entity_name.lower().strip()

            # Verifikasi Entitas
            if entity_lower not in content_combined:
                continue

            # Verifikasi Kata Kunci Sesuai Mode
            if keyword_list:
                if match_operator == "AND":
                    # Wajib memuat seluruh kata kunci
                    if not all(kw.lower() in content_combined for kw in keyword_list):
                        continue
                else:
                    # Wajib memuat minimal 1 kata kunci yang dimaksud
                    if not any(kw.lower() in content_combined for kw in keyword_list):
                        continue

            # Resolving Link Asli Penerbit
            clean_link = resolve_clean_url(raw_link)
            if clean_link in seen_links:
                continue
            seen_links.add(clean_link)

            # Filter Situs Pemerintah jika lolos
            if exclude_gov and (".go.id" in clean_link.lower() or ".go.id" in source.lower()):
                continue

            lower_link = clean_link.lower()
            if channel_type == "medsos" or any(s in lower_link for s in ["instagram.com", "facebook.com", "threads.net", "twitter.com", "x.com", "tiktok.com"]):
                channel_type = "medsos"
                if "instagram.com" in lower_link:
                    source = "Instagram"
                elif "facebook.com" in lower_link:
                    source = "Facebook"
                elif "threads.net" in lower_link:
                    source = "Threads"
                elif "twitter.com" in lower_link or "x.com" in lower_link:
                    source = "X (Twitter)"
                elif "tiktok.com" in lower_link:
                    source = "TikTok"
            elif channel_type == "local_web":
                source = f"{source} (Web/Lokal)"

            bullet_points = extract_bullet_summary(clean_title, snippet, min_bullets=2)

            t_low = content_combined
            neg_count = sum(1 for w in negative_lexicon if w in t_low)
            pos_count = sum(1 for w in positive_lexicon if w in t_low)

            if neg_count > pos_count:
                sentiment = "Negatif"
                pr_tone = "High Risk / Potential Crisis"
                if channel_type == "medsos":
                    pr_rec = (
                        "Rekomendasi Respons Medsos:\n"
                        "1. Aktifkan Social Media Response Unit; berikan tanggapan penyejuk di kolom komentar dalam 1 jam.\n"
                        "2. Jangan menghapus kritik warganet secara sepihak; arahkan ke formulir pelaporan resmi.\n"
                        "3. Siapkan FAQ untuk menangkal rumor viral yang belum tentu faktual."
                    )
                else:
                    pr_rec = (
                        "Rekomendasi Konsultan Media:\n"
                        "1. Segera siapkan 'Holding Statement' resmi dan sebarkan hak jawab proporsional.\n"
                        "2. Adakan media briefing proaktif kepada jurnalis peliput untuk meluruskan fakta.\n"
                        "3. Siapkan lembar FAQ mitigasi isu bagi juru bicara institusi."
                    )
            elif pos_count > neg_count:
                sentiment = "Positif"
                pr_tone = "Strategic Traction / Good PR"
                if channel_type == "medsos":
                    pr_rec = (
                        "Rekomendasi Respons Medsos:\n"
                        "1. Repost dan amplifikasi postingan ke kanal media sosial resmi institusi.\n"
                        "2. Jalin engagement dengan memberikan likes/komentar apresiasi resmi.\n"
                        "3. Gunakan materi sebagai Social Proof kredibilitas institusi."
                    )
                else:
                    pr_rec = (
                        "Rekomendasi Konsultan Media:\n"
                        "1. Amplifikasi narasi positif melalui siaran pers lanjutan ke media nasional.\n"
                        "2. Jadikan rujukan pada public expose pimpinan dan bahan laporan berkala.\n"
                        "3. Jaga hubungan relasi media dengan jurnalis peliput."
                    )
            else:
                sentiment = "Netral"
                pr_tone = "Informative / General Update"
                pr_rec = (
                    "Rekomendasi Konsultan Media:\n"
                    "1. Monitor dinamika pemberitaan lanjutan (follow-up coverage).\n"
                    "2. Sediakan factsheet ringkas jika terdapat awak media yang meminta rincian data."
                )

            card_bytes = generate_press_capture_card(
                channel_type, source, clean_title, published, snippet, sentiment
            )

            items.append({
                "channel_type": channel_type,
                "title": clean_title,
                "source": source,
                "published": published,
                "link": clean_link,
                "snippet": snippet,
                "bullets": bullet_points,
                "sentiment": sentiment,
                "pr_tone": pr_tone,
                "pr_rec": pr_rec,
                "image_bytes": card_bytes
            })

            if max_results != "unlimited" and len(items) >= int(max_results):
                break
        if max_results != "unlimited" and len(items) >= int(max_results):
            break

    return items, advanced_q

# ==========================================
# 6. GENERATOR SLIDE PPTX
# ==========================================
def generate_pptx_deck(entity_name, org_type, date_period_label, articles):
    prs = Presentation()
    prs.slide_width = Inches(13.33)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    # COVER SLIDE
    cover_slide = prs.slides.add_slide(blank_layout)
    title_box = cover_slide.shapes.add_textbox(Inches(1.0), Inches(1.8), Inches(11.33), Inches(4.0))
    tf = title_box.text_frame
    tf.word_wrap = True
    
    p0 = tf.paragraphs[0]
    p0.text = "MEDIA & SOCIAL INTELLIGENCE AUDIT"
    p0.font.size = Pt(36)
    p0.font.bold = True
    p0.font.color.rgb = RGBColor(20, 35, 60)

    p1 = tf.add_paragraph()
    p1.text = f"Target Entitas: {entity_name} | Kategori: {org_type}"
    p1.font.size = Pt(22)
    p1.font.color.rgb = RGBColor(70, 80, 95)

    p2 = tf.add_paragraph()
    p2.text = f"Periode Pemantauan: {date_period_label}"
    p2.font.size = Pt(16)
    p2.font.bold = True
    p2.font.color.rgb = RGBColor(14, 116, 144)

    p3 = tf.add_paragraph()
    p3.text = "Disusun oleh: Senior Strategic Media & Reputation Advisory"
    p3.font.size = Pt(13)
    p3.font.italic = True
    p3.font.color.rgb = RGBColor(120, 130, 145)

    # CONTENT SLIDES
    for idx, art in enumerate(articles, 1):
        slide = prs.slides.add_slide(blank_layout)

        header_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.7), Inches(0.6))
        htf = header_box.text_frame
        hp = htf.paragraphs[0]
        channel_tag = "MEDSOS" if art['channel_type'] == "medsos" else "MEDIA"
        hp.text = f"DOKUMENTASI #{idx} [{channel_tag}] : {art['source'].upper()}"
        hp.font.size = Pt(17)
        hp.font.bold = True
        hp.font.color.rgb = RGBColor(30, 41, 59)

        if art.get("image_bytes"):
            img_stream = io.BytesIO(art["image_bytes"])
            slide.shapes.add_picture(img_stream, Inches(0.8), Inches(1.2), width=Inches(5.6), height=Inches(4.8))

        # TAUTAN ASLI DI SLIDE
        link_box = slide.shapes.add_textbox(Inches(0.8), Inches(6.1), Inches(5.6), Inches(0.8))
        ltf = link_box.text_frame
        ltf.word_wrap = True
        lp = ltf.paragraphs[0]
        lp.text = f"🔗 Tautan Sumber Penerbit ({art['source']}):"
        lp.font.size = Pt(10)
        lp.font.bold = True
        lp.font.color.rgb = RGBColor(71, 85, 105)

        link_p = ltf.add_paragraph()
        run = link_p.add_run()
        run.text = art['link'][:65] + ("..." if len(art['link']) > 65 else "")
        run.font.size = Pt(10)
        run.font.underline = True
        run.font.color.rgb = RGBColor(37, 99, 235)
        run.hyperlink.address = art['link']

        # RINGKASAN & ADVISORY
        right_box = slide.shapes.add_textbox(Inches(6.7), Inches(1.2), Inches(5.8), Inches(5.7))
        rtf = right_box.text_frame
        rtf.word_wrap = True

        rp1 = rtf.paragraphs[0]
        rp1.text = "RINGKASAN PEMBERITAAN (KEY TAKEAWAYS)"
        rp1.font.size = Pt(11)
        rp1.font.bold = True
        rp1.font.color.rgb = RGBColor(14, 116, 144)

        for b in art['bullets']:
            bp = rtf.add_paragraph()
            bp.text = f"• {b}"
            bp.font.size = Pt(11)
            bp.font.color.rgb = RGBColor(30, 41, 59)

        div_p = rtf.add_paragraph()
        div_p.text = "―" * 28
        div_p.font.size = Pt(9)
        div_p.font.color.rgb = RGBColor(203, 213, 225)

        adv_title = rtf.add_paragraph()
        adv_title.text = "STRATEGIC MEDIA & SOCIAL ADVISORY"
        adv_title.font.size = Pt(11)
        adv_title.font.bold = True
        adv_title.font.color.rgb = RGBColor(14, 116, 144)

        sent_p = rtf.add_paragraph()
        sent_p.text = f"Sentimen: {art['sentiment'].upper()} ({art['pr_tone']})"
        sent_p.font.size = Pt(12)
        sent_p.font.bold = True
        if art['sentiment'] == "Negatif":
            sent_p.font.color.rgb = RGBColor(220, 38, 38)
        elif art['sentiment'] == "Positif":
            sent_p.font.color.rgb = RGBColor(22, 163, 74)
        else:
            sent_p.font.color.rgb = RGBColor(202, 138, 4)

        rec_p = rtf.add_paragraph()
        rec_p.text = f"{art['pr_rec']}"
        rec_p.font.size = Pt(11)
        rec_p.font.color.rgb = RGBColor(15, 23, 42)

    # EXECUTIVE SUMMARY SLIDE
    summary_slide = prs.slides.add_slide(blank_layout)
    sum_header = summary_slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.7), Inches(0.8))
    stf = sum_header.text_frame
    sp = stf.paragraphs[0]
    sp.text = "EXECUTIVE SUMMARY: STRATEGIC ACTIONABLE RECOMMENDATIONS"
    sp.font.size = Pt(20)
    sp.font.bold = True
    sp.font.color.rgb = RGBColor(20, 35, 60)

    col_w, gap, start_x, top_y, card_h = Inches(3.7), Inches(0.3), Inches(0.8), Inches(1.3), Inches(5.6)

    # Negatif
    neg_box = summary_slide.shapes.add_textbox(start_x, top_y, col_w, card_h)
    ntf = neg_box.text_frame
    ntf.word_wrap = True
    np1 = ntf.paragraphs[0]
    np1.text = "ISU NEGATIF / KRISIS"
    np1.font.size = Pt(14)
    np1.font.bold = True
    np1.font.color.rgb = RGBColor(220, 38, 38)
    np2 = ntf.add_paragraph()
    np2.text = (
        "Langkah Mitigasi Eksekutif:\n"
        "• Segera rilis Hak Jawab / Holding Statement resmi dalam 24 jam.\n"
        "• Lakukan penelusuran fakta objektif internal sebelum juru bicara berbicara di publik.\n"
        "• Berikan penanganan santun & terarah pada kanal keluhan publik/medsos.\n"
        "• Pantau agar isu lokal tidak tereskalasi menjadi isu nasional di media tier-1."
    )
    np2.font.size = Pt(11)
    np2.font.color.rgb = RGBColor(51, 65, 85)

    # Netral
    neu_box = summary_slide.shapes.add_textbox(start_x + col_w + gap, top_y, col_w, card_h)
    netf = neu_box.text_frame
    netf.word_wrap = True
    nep1 = netf.paragraphs[0]
    nep1.text = "ISU NETRAL / INFORMATIF"
    nep1.font.size = Pt(14)
    nep1.font.bold = True
    nep1.font.color.rgb = RGBColor(202, 138, 4)
    nep2 = netf.add_paragraph()
    nep2.text = (
        "Langkah Pengawasan & Edukasi:\n"
        "• Cermati perkembangan isu apakah terdapat potensi pergeseran persepsi publik.\n"
        "• Sediakan lembar fakta (factsheet) pendukung bagi wartawan yang meminta data komparasi.\n"
        "• Edukasi publik melalui artikel opini atau infografis sederhana di situs web resmi."
    )
    nep2.font.size = Pt(11)
    nep2.font.color.rgb = RGBColor(51, 65, 85)

    # Positif
    pos_box = summary_slide.shapes.add_textbox(start_x + (col_w + gap) * 2, top_y, col_w, card_h)
    ptf = pos_box.text_frame
    ptf.word_wrap = True
    pp1 = ptf.paragraphs[0]
    pp1.text = "ISU POSITIF / PRESTASI"
    pp1.font.size = Pt(14)
    pp1.font.bold = True
    pp1.font.color.rgb = RGBColor(22, 163, 74)
    pp2 = ptf.add_paragraph()
    pp2.text = (
        "Langkah Amplifikasi Reputasi:\n"
        "• Perluas jangkauan narasi prestasi ke owned media (LinkedIn, Press Release, IG Reels/Threads).\n"
        "• Manfaatkan momentum keberhasilan untuk Thought Leadership Direksi/Pimpinan.\n"
        "• Jalin silaturahmi apresiasi dengan jurnalis yang memberitakan secara berimbang."
    )
    pp2.font.size = Pt(11)
    pp2.font.color.rgb = RGBColor(51, 65, 85)

    output_stream = io.BytesIO()
    prs.save(output_stream)
    output_stream.seek(0)
    return output_stream

# ==========================================
# 7. INTERFACE STREAMLIT
# ==========================================
with st.sidebar:
    st.header("⚙️ Konfigurasi Monitoring")
    entity_input = st.text_input("Nama Entitas / Target", value="Bank Mandiri")
    org_type_input = st.selectbox("Kategori Entitas", [
        "Perbankan / Finansial", 
        "BUMN / Pemerintahan", 
        "Kehutanan / Lingkungan Hidup",
        "Korporasi Swasta", 
        "Lembaga Publik / NGO", 
        "Personal / Tokoh Publik"
    ])

    st.subheader("🎯 Kata Kunci Isu Spesifik")
    keywords_input = st.text_area(
        "Daftar Topik / Isu (Pisahkan dengan Koma):",
        value='csr, taspen, sosialisasi, daerah',
        help="Contoh: csr, taspen, sosialisasi, daerah"
    )

    # PILIHAN LOGIKA BOOLEAN KATA KUNCI
    match_op = st.radio(
        "Kombinasi Logika Kata Kunci:",
        options=["AND (Wajib Ada Semua Kata Kunci)", "OR (Cukup Salah Satu Saja)"],
        index=0,
        help="Pilih 'AND' agar hasil penjaringan benar-benar fokus ke topik Taspen & CSR bersamaan, dan tidak melebar ke berita lain."
    )
    selected_operator = "AND" if "AND" in match_op else "OR"

    match_title_only = st.checkbox("Wajibkan Kata Kunci di Judul", value=False)

    st.subheader("🛡️ Filter Anti-Bias (Pengecualian Media)")
    exclude_gov_media = st.checkbox("Kecualikan Media Pemerintah (*.go.id)", value=True)
    custom_exclude = st.text_input("Kecualikan Domain Resmi Target:", value="bankmandiri.co.id, mandiri.co.id")

    st.subheader("🌐 Jangkauan Sumber Media")
    inc_news = st.checkbox("Media Nasional (Tier-1)", value=True)
    inc_small = st.checkbox("Media Lokal & Website Kecil", value=True)
    inc_social = st.checkbox("Media Sosial (IG, FB, Threads, X, TikTok)", value=True)

    st.subheader("⏳ Rentang Waktu Pemantauan")
    time_options = {
        "1 Jam Terakhir": "1h",
        "3 Jam Terakhir": "3h",
        "6 Jam Terakhir": "6h",
        "12 Jam Terakhir": "12h",
        "24 Jam Terakhir (1 Hari)": "1d",
        "2 Hari Terakhir": "2d",
        "3 Hari Terakhir": "3d",
        "7 Hari Terakhir (1 Minggu)": "7d",
        "30 Hari Terakhir (1 Bulan)": "30d",
        "Semua Waktu (Tanpa Batas)": "all"
    }
    selected_time_label = st.selectbox("Pilih Rentang Waktu", list(time_options.keys()), index=4)
    time_val = time_options[selected_time_label]
    calculated_date_label = calculate_date_range_label(time_val)
    st.caption(f"🗓️ **Periode Aktif:** {calculated_date_label}")

    st.subheader("📊 Kuota Temuan")
    limit_mode = st.radio("Mode Kuota:", ["Tentukan Jumlah", "Bebas (Seluruhnya)"])
    max_news_count = st.number_input("Jumlah Temuan", min_value=1, max_value=100, value=15, step=5) if limit_mode == "Tentukan Jumlah" else "unlimited"

    run_btn = st.button("🚀 Mulai Pemindaian Presisi", type="primary", use_container_width=True)

# ==========================================
# 8. HASIL PEMINDAIAN & EKSPOR
# ==========================================
if run_btn:
    with st.spinner("Memindai berita, memverifikasi keterkaitan kata kunci, dan membersihkan tautan..."):
        items, built_query = fetch_media_and_social(
            entity_name=entity_input,
            raw_keywords=keywords_input,
            time_window=time_val,
            max_results=max_news_count,
            include_news=inc_news,
            include_social=inc_social,
            include_small_media=inc_small,
            match_operator=selected_operator,
            match_in_title=match_title_only,
            exclude_gov=exclude_gov_media,
            custom_exclude=custom_exclude
        )
        st.session_state['news_data'] = items
        st.session_state['entity'] = entity_input
        st.session_state['org_type'] = org_type_input
        st.session_state['date_label'] = calculated_date_label
        st.session_state['last_query'] = built_query

if 'news_data' in st.session_state:
    articles = st.session_state.get('news_data', [])
    if 'last_query' in st.session_state:
        st.caption(f"🔍 **Query Penjaringan Presisi:** `{st.session_state['last_query']}`")

    if not articles:
        st.warning("⚠️ Tidak ditemukan berita yang memenuhi seluruh kombinasi kata kunci tersebut. Coba ganti opsi logika menjadi 'OR' atau kurangi kata kunci yang terlalu spesifik.")
    else:
        total_items = len(articles)
        pos_count = sum(1 for a in articles if a['sentiment'] == "Positif")
        neg_count = sum(1 for a in articles if a['sentiment'] == "Negatif")
        neu_count = sum(1 for a in articles if a['sentiment'] == "Netral")

        c1, c2, c3, c4 = st.columns(4)
        c1.metric("Total Temuan Terverifikasi", total_items)
        c2.metric("Sentimen Positif", pos_count)
        c3.metric("Sentimen Negatif (Krisis)", neg_count)
        c4.metric("Sentimen Netral", neu_count)

        st.divider()

        pptx_file = generate_pptx_deck(
            st.session_state['entity'], 
            st.session_state['org_type'], 
            st.session_state['date_label'],
            articles
        )
        st.download_button(
            label=f"📥 Unduh Laporan Slide Deck ({total_items + 2} Slide .PPTX Siap Paparan)",
            data=pptx_file,
            file_name=f"PR_Audit_{st.session_state['entity'].replace(' ', '_')}.pptx",
            mime="application/vnd.openxmlformats-officedocument.presentationml.presentation"
        )

        st.subheader("📋 Kliping Digital & Rekomendasi Media Konsultan")
        for idx, art in enumerate(articles, 1):
            with st.container():
                col_left, col_right = st.columns([1, 1], gap="medium")
                with col_left:
                    st.image(art["image_bytes"], caption=f"Capture #{idx} - {art['source']}", use_container_width=True)
                    st.markdown(f"🔗 **Tautan Langsung:** [{art['link'][:70]}...]({art['link']})")
                with col_right:
                    st.markdown(f"### #{idx}: {art['title']}")
                    st.caption(f"**Sumber Penerbit:** {art['source']} | **Waktu:** {art['published']}")
                    st.markdown("#### 📝 Ringkasan Isi Berita (Key Points):")
                    for b in art['bullets']:
                        st.markdown(f"- {b}")
                    st.markdown("---")
                    badge_color = "🟢" if art['sentiment'] == "Positif" else "🔴" if art['sentiment'] == "Negatif" else "🟡"
                    st.markdown(f"#### {badge_color} Sentimen: **{art['sentiment']}** ({art['pr_tone']})")
                    st.info(art['pr_rec'])
                st.markdown("---")
